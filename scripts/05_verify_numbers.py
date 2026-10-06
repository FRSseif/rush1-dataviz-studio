"""Phase 8 — Vérifie que chaque chiffre de la Synthèse correspond au calcul Python.

1. Recalcule les chiffres clés DIRECTEMENT depuis les CSV bruts (nettoyage refait en
   mémoire, statistiques avec scipy), indépendamment du classeur.
2. Compare ces valeurs aux chiffres du classeur (valeurs enregistrées).
3. Vérifie que les nombres cités dans les phrases de la Synthèse sont les bons.
4. Avec --excel : fait recalculer TOUT le classeur par Microsoft Excel et compare
   chaque cellule numérique au calcul Python (aucune erreur, aucun écart toléré).

Usage : .venv/bin/python scripts/05_verify_numbers.py [--excel] [chemin.xlsx]
"""
import importlib
import math
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
clean = importlib.import_module("01_clean")
build = importlib.import_module("03_build_workbook")
CLIENT = ROOT / "deliverable" / "Dataviz_Studio_TikTok_Client.xlsx"
fr, pct, pval = build.fr, build.pct, build.pval


def donnees_brutes():
    tv = clean.load_video_extract("trending_videos.csv", "Tendances")
    lk = clean.load_video_extract("tiktok_collected_liked_videos.csv", "Likées par @tiktok")
    fh = clean.load_funny()
    return clean.derive(pd.concat([tv, fh, lk], ignore_index=True))


def chiffres_cles(df):
    """Chiffres de la Synthèse, recalculés sans passer par le classeur."""
    t = df[df["echantillon"] == "Tendances"]
    h = df[df["echantillon"] == "Humour"]
    petit = h[h["abonnes"] < 1e6]
    gros = h[h["abonnes"] >= 8e6]

    def ecart(d, mask):
        return (d.loc[mask, "engagement"].median() - d.loc[~mask, "engagement"].median()) * 100

    themes = t.groupby("theme")["engagement"].median()
    return {
        "rho_abonnes_vues": stats.spearmanr(h["abonnes"], h["vues"])[0],
        "rho_abonnes_engagement": stats.spearmanr(h["abonnes"], h["engagement"])[0],
        "eng_moins_1M": petit["engagement"].median(),
        "eng_plus_8M": gros["engagement"].median(),
        "top_moins_1M": petit["top_performer"].mean(),
        "top_plus_8M": gros["top_performer"].mean(),
        "ratio_abonnes": gros["abonnes"].median() / petit["abonnes"].median(),
        "ratio_vues": gros["vues"].median() / petit["vues"].median(),
        "ratio_vues_par_abonne": (petit["vues"] / petit["abonnes"]).median()
                                 / (gros["vues"] / gros["abonnes"]).median(),
        "ecart_fyp": ecart(t, t["hashtag_generique"].astype(bool)),
        "p_fyp": stats.mannwhitneyu(t.loc[t["hashtag_generique"].astype(bool), "engagement"],
                                    t.loc[~t["hashtag_generique"].astype(bool), "engagement"],
                                    alternative="two-sided").pvalue,
        "ecart_duree": ecart(t, t["duree_s"] > 15),
        "eng_plus_15s": t.loc[t["duree_s"] > 15, "engagement"].median(),
        "eng_15s_ou_moins": t.loc[t["duree_s"] <= 15, "engagement"].median(),
        "top_plus_15s": t.loc[t["duree_s"] > 15, "top_performer"].mean(),
        "top_15s_ou_moins": t.loc[t["duree_s"] <= 15, "top_performer"].mean(),
        "p_duree": stats.mannwhitneyu(t.loc[t["duree_s"] > 15, "engagement"],
                                      t.loc[t["duree_s"] <= 15, "engagement"],
                                      alternative="two-sided").pvalue,
        "ecart_theme": (themes.max() - themes.min()) * 100,
        "p_theme": stats.kruskal(*[g["engagement"] for _, g in t.groupby("theme")]).pvalue,
        "seuil_tendances": t["engagement"].quantile(0.75),
        "seuil_humour": h["engagement"].quantile(0.75),
    }


class Rapport:
    def __init__(self):
        self.lignes = []

    def check(self, nom, ok, detail=""):
        self.lignes.append((nom, bool(ok), detail))

    def afficher(self):
        w = max(len(n) for n, _, _ in self.lignes)
        for nom, ok, detail in self.lignes:
            print(f"{'PASS' if ok else 'FAIL'}  {nom:<{w}}  {detail}")
        print(f"\n{sum(ok for _, ok, _ in self.lignes)}/{len(self.lignes)} contrôles PASS")
        return all(ok for _, ok, _ in self.lignes)


def proche(a, b, tol=1e-9):
    return isinstance(b, (int, float)) and math.isclose(a, b, rel_tol=tol, abs_tol=tol)


def main(path, avec_excel):
    k = chiffres_cles(donnees_brutes())
    wb = openpyxl.load_workbook(path, data_only=True)
    syn, gc, me = wb["Synthèse"], wb["Gros comptes"], wb["Méthode & limites"]
    R = Rapport()

    # 2. Chiffres clés de la Synthèse (colonne D, lignes 6 à 10).
    attendus = [("D6", "rho_abonnes_vues", "ρ abonnés ↔ vues"),
                ("D7", "eng_moins_1M", "engagement médian < 1 M d'abonnés"),
                ("D8", "ecart_fyp", "écart #fyp (points, tendances)"),
                ("D9", "ecart_duree", "écart durée > 15 s (points, tendances)"),
                ("D10", "ecart_theme", "écart max entre sujets (points, tendances)")]
    for cell, cle, lib in attendus:
        v = syn[cell].value
        R.check(f"Synthèse {cell} : {lib}", proche(k[cle], v),
                f"classeur {v!r} / Python {k[cle]!r}")
    R.check("Gros comptes I7 : rapport vues par abonné", proche(k["ratio_vues_par_abonne"],
                                                               gc["I7"].value),
            f"{gc['I7'].value!r} / {k['ratio_vues_par_abonne']!r}")
    R.check("Seuils top performers (Méthode C9, C11)",
            proche(k["seuil_tendances"], me["C9"].value) and proche(k["seuil_humour"],
                                                                    me["C11"].value),
            f"{me['C9'].value:.6f} / {me['C11'].value:.6f}")

    # 3. Nombres cités dans les phrases de la Synthèse.
    texte = " ".join(str(c.value) for row in syn.iter_rows() for c in row
                     if isinstance(c.value, str))
    cites = {
        "36 fois plus d'abonnés": f"{fr(k['ratio_abonnes'], 0)} fois plus d'abonnés",
        "1,3 fois plus de vues": f"{fr(k['ratio_vues'], 1)} fois plus de vues",
        "engagement > 8 M": pct(k["eng_plus_8M"]),
        "top performers < 1 M": f"{pct(k['top_moins_1M'], 0)} de top performers",
        "top performers > 8 M": f"contre {pct(k['top_plus_8M'], 0)})",
        "p du test #fyp": f"p = {pval(k['p_fyp'])}",
        "engagement > 15 s": f"{pct(k['eng_plus_15s'])} d'engagement",
        "engagement ≤ 15 s": f"contre {pct(k['eng_15s_ou_moins'])}",
        "top performers > 15 s": f"{pct(k['top_plus_15s'], 0)} de top",
        "top performers ≤ 15 s": f"contre {pct(k['top_15s_ou_moins'], 0)}, p = {pval(k['p_duree'])}",
        "p du test des sujets": f"p = {pval(k['p_theme'])}",
        "vues par abonné (fois)": f"{fr(k['ratio_vues_par_abonne'], 0)} fois plus de vues par",
    }
    for lib, attendu in cites.items():
        R.check(f"Texte : {lib}", attendu in texte, f"« {attendu} »")

    # 4. Recalcul complet par Excel.
    if avec_excel:
        from excel_mac import roundtrip
        dest = ROOT / "work" / "controle_excel.xlsx"
        info = roundtrip(Path(path), dest, ["Facteurs", "Gros comptes", "Catégories & TreeMap"])
        R.check("Excel ouvre le fichier sans réparation", info.get("feuilles") == "6",
                str(info))
        R.check("Excel reconnaît les 3 TCD", all(info.get(f) == "1" for f in
                                                 ("Facteurs", "Gros comptes",
                                                  "Catégories & TreeMap")))
        xl = openpyxl.load_workbook(dest, data_only=True)
        n, ecarts, erreurs = 0, [], []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    vx = xl[ws.title][c.coordinate].value
                    if isinstance(vx, str) and vx.startswith("#"):
                        erreurs.append(f"{ws.title}!{c.coordinate}")
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        n += 1
                        if not proche(c.value, vx):
                            ecarts.append(f"{ws.title}!{c.coordinate}")
        R.check("Excel : aucune cellule en erreur après recalcul", not erreurs,
                ", ".join(erreurs[:5]))
        R.check("Excel = Python sur toutes les cellules numériques", not ecarts,
                f"{n} cellules comparées" + (f", écarts : {ecarts[:5]}" if ecarts else ""))
    return R.afficher()


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    ok = main(args[0] if args else CLIENT, "--excel" in sys.argv)
    sys.exit(0 if ok else 1)
